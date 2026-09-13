
import io
import base64
import json
import os
import re
from typing import Any, Dict, List
from pathlib import Path

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, Mm
from docx.oxml.ns import qn
from google import genai
from google.genai import types
from io import BytesIO


TEMPLATE_PATH = Path(__file__).with_name("school_learning_plan_template.docx")


# ============================================================
# APP CONFIG
# ============================================================

st.set_page_config(
    page_title="Daily Lesson Plan Generator",
    page_icon="",
    layout="centered",
)

LOGO_PATH = Path(__file__).with_name("smaj_logo.png")
try:
    LOGO_B64 = base64.b64encode(LOGO_PATH.read_bytes()).decode("ascii")
except Exception:
    LOGO_B64 = ""

BG_PATH = Path(__file__).with_name("smaj_blue_background.png")
try:
    BG_B64 = base64.b64encode(BG_PATH.read_bytes()).decode("ascii")
except Exception:
    BG_B64 = ""

st.markdown(f"""
<style>
:root {{
    --blue-deep: #061A40;
    --blue-main: #0B3D91;
    --blue-mid: #1368CE;
    --blue-bright: #2F80ED;
    --blue-soft: #EAF4FF;
    --blue-pale: #F7FBFF;
    --white: #FFFFFF;
    --ink: #08224D;
}}

/* User-provided technology background + dark overlay for readable text */
.stApp {{
    background:
        linear-gradient(rgba(3, 18, 51, .42), rgba(5, 36, 83, .52)),
        url("data:image/png;base64,{BG_B64}") center center / cover fixed no-repeat;
    color: #FFFFFF;
}}

.block-container {{
    max-width: 960px;
    padding-top: 1.35rem;
    padding-bottom: 3rem;
}}

/* Main text remains readable on blue background */
h1, h2, h3,
div[data-testid="stWidgetLabel"] p {{
    color: #FFFFFF !important;
    text-shadow: 0 1px 2px rgba(0,0,0,.22);
}}

/* Hero */
.hero-card {{
    background: rgba(255,255,255,.97);
    color: var(--ink);
    border-radius: 24px;
    padding: 1.4rem 1.25rem 1.15rem;
    text-align: center;
    margin-bottom: 1.4rem;
    box-shadow: 0 18px 44px rgba(0, 18, 55, .38);
    border: 1px solid rgba(255,255,255,.86);
    backdrop-filter: blur(8px);
}}
.hero-card img {{
    width: 116px;
    height: 116px;
    object-fit: contain;
    margin-bottom: .35rem;
}}
.hero-card h1 {{
    color: var(--blue-main) !important;
    text-shadow: none !important;
    margin: .08rem 0 .22rem;
    font-size: 2.08rem;
    font-weight: 850;
}}
.hero-card p {{
    color: #244A7C !important;
    margin: 0;
    font-size: .98rem;
}}
.hero-card .school-name {{
    color: var(--blue-deep) !important;
    font-size: .9rem;
    font-weight: 800;
    letter-spacing: .035em;
    margin-bottom: .18rem;
}}
.hero-card .developer-line {{
    margin-top: .45rem;
    color: #376CA8 !important;
    font-size: .8rem;
    font-weight: 700;
}}

/* Section cards */
div[data-testid="stVerticalBlock"] > div:has(> div[data-testid="stFileUploader"]) {{
    border-radius: 16px;
}}

/* Input labels and captions */
label, .stCaption, [data-testid="stCaptionContainer"] p {{
    color: #F5FAFF !important;
}}
.small {{
    color: #D8EBFF !important;
}}

/* White input fields */
div[data-testid="stTextInput"] input,
div[data-testid="stTextArea"] textarea,
div[data-baseweb="select"] > div {{
    background: rgba(255,255,255,.98) !important;
    color: #08224D !important;
    border: 2px solid #73B7FF !important;
    border-radius: 11px !important;
    box-shadow: 0 4px 12px rgba(0,40,100,.12);
}}
div[data-testid="stTextInput"] input::placeholder,
div[data-testid="stTextArea"] textarea::placeholder {{
    color: #6C84A6 !important;
    opacity: 1 !important;
}}
div[data-baseweb="select"] span,
div[data-baseweb="select"] div {{
    color: #08224D !important;
}}


/* Extra visibility for dropdown text/options */
div[data-baseweb="select"] > div {{
    min-height: 3rem !important;
    font-size: 1rem !important;
    font-weight: 750 !important;
}}
div[data-baseweb="select"] svg {{
    fill: #08224D !important;
}}
ul[role="listbox"],
div[role="listbox"] {{
    background: #FFFFFF !important;
}}
li[role="option"],
div[role="option"] {{
    background: #FFFFFF !important;
    color: #08224D !important;
    font-size: 1rem !important;
    font-weight: 700 !important;
}}
li[role="option"]:hover,
div[role="option"]:hover {{
    background: #EAF4FF !important;
    color: #062D67 !important;
}}
div[data-testid="stWidgetLabel"] p {{
    color: #FFFFFF !important;
    font-weight: 800 !important;
    font-size: 1rem !important;
}}

/* Upload boxes */
div[data-testid="stFileUploader"] {{
    background: rgba(9, 56, 116, .62);
    border: 1px solid rgba(255,255,255,.26);
    border-radius: 16px;
    padding: .4rem;
    box-shadow: 0 8px 22px rgba(0,18,55,.18);
}}
div[data-testid="stFileUploaderDropzone"] {{
    background: rgba(255,255,255,.98) !important;
    border: 2px dashed #3D9CFF !important;
    border-radius: 14px !important;
}}
div[data-testid="stFileUploaderDropzone"] * {{
    color: #0B3D91 !important;
}}

/* Buttons */
.stButton > button, .stDownloadButton > button {{
    width: 100%;
    min-height: 3rem;
    border-radius: 12px;
    font-weight: 850;
    background: linear-gradient(135deg, #0B3D91, #1368CE) !important;
    color: #FFFFFF !important;
    border: 2px solid #70B7FF !important;
    box-shadow: 0 10px 24px rgba(0, 28, 82, .30);
}}
.stButton > button:hover, .stDownloadButton > button:hover {{
    background: linear-gradient(135deg, #1368CE, #2F80ED) !important;
    color: #FFFFFF !important;
    border-color: #B8DCFF !important;
    transform: translateY(-1px);
}}

/* Requirement status cards */
.status-ok {{
    padding: .62rem .82rem;
    border-radius: 10px;
    background: rgba(255,255,255,.97) !important;
    border-left: 5px solid #2F80ED !important;
    color: #0B3D91 !important;
    margin-bottom: .45rem;
    box-shadow: 0 4px 12px rgba(0,25,75,.14);
}}
.status-no {{
    padding: .62rem .82rem;
    border-radius: 10px;
    background: #FFF0F0 !important;
    border-left: 5px solid #D64545 !important;
    color: #8B1E2D !important;
    margin-bottom: .45rem;
}}

/* Alerts */
div[data-testid="stAlert"] {{
    border-radius: 12px;
}}
div[data-testid="stAlert"] * {{
    color: inherit !important;
}}

/* Divider */
hr {{
    border-color: rgba(255,255,255,.32) !important;
}}

/* Footer */
.developer {{
    text-align: center;
    color: #E6F3FF !important;
    font-size: .88rem;
    margin-top: 1.2rem;
    font-weight: 700;
}}
</style>
""", unsafe_allow_html=True)

logo_html = f'<img src="data:image/png;base64,{LOGO_B64}" alt="SMAJ Logo">' if LOGO_B64 else ''
st.markdown(
    f"""
    <div class="hero-card">
        {logo_html}
        <div class="school-name">ST. MARY'S ACADEMY OF JASAAN, INC.</div>
        <h1>SMAJ TeachPlan</h1>
        <p>Daily Lesson Plan Generator</p>
        <p style="margin-top:.35rem;">Simple, aligned, student-friendly lesson planning for Marian educators.</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SECRET / ENV SETTINGS
# ============================================================

def get_secret(name: str, default: str = "") -> str:
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass
    return os.getenv(name, default)

GEMINI_API_KEY = get_secret("GEMINI_API_KEY")
GEMINI_MODEL = get_secret("GEMINI_MODEL", "gemini-2.5-flash-lite")
SITE_PASSWORD = get_secret("SITE_PASSWORD", "")


# ============================================================
# OPTIONAL SITE PASSWORD
# ============================================================

if SITE_PASSWORD:
    if "site_unlocked" not in st.session_state:
        st.session_state.site_unlocked = False

    if not st.session_state.site_unlocked:
        st.subheader("Teacher Access")
        pw = st.text_input("Enter site password", type="password")
        if st.button("Open Generator"):
            if pw == SITE_PASSWORD:
                st.session_state.site_unlocked = True
                st.rerun()
            else:
                st.error("Incorrect password.")
        st.stop()


# ============================================================
# DOCX INPUT READER
# ============================================================

def read_docx(uploaded_file) -> str:
    doc = Document(uploaded_file)
    parts: List[str] = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = " ".join(
                    p.text.strip()
                    for p in cell.paragraphs
                    if p.text.strip()
                )
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                parts.append(" | ".join(row_text))

    return "\n".join(parts)


# ============================================================
# STRUCTURED OUTPUT SCHEMA
# ============================================================

SCHEMA: Dict[str, Any] = {
    "school_name": "",
    "school_year": "",
    "subject": "",
    "grade_level": "",
    "term": "",
    "session": "",
    "topic": "",
    "unit": "",
    "date": "",
    "learning_competencies": [""],
    "transfer_goal": "",
    "essential_understanding": "",
    "essential_question": "",
    "preliminaries": {
        "review": {"lots_question": "", "hots_question": ""},
        "focus": "",
        "resources": [""],
        "motivation": {
            "title": "",
            "instruction": "",
            "content": "",
            "guide_questions": [""]
        },
        "activating_prior_knowledge": ""
    },
    "lesson_development": {
        "presentation_of_concept": [""],
        "activities": [
            {
                "title": "",
                "type": "",
                "instruction": "",
                "steps": [""],
                "guide_questions": [""]
            }
        ],
        "broadening_of_concept": {
            "leading_question": "",
            "exploring_question": "",
            "connecting_question": "",
            "essential_question": ""
        },
        "integration": {
            "ignacian_core_value": {"name": "", "connection": ""},
            "related_core_value": {"name": "", "connection": ""},
            "social_orientation": {"name": "", "question_or_connection": ""},
            "lesson_across_discipline": {"name": "", "question_or_connection": ""},
            "biblical_text_reflection": {"reference": "", "text": "", "reflection": ""}
        }
    },
    "evaluation_assessment": {
        "formative": [{"name": "", "instruction": "", "items": [""]}],
        "summative": [{"name": "", "instruction": "", "items": [""]}]
    },
    "summary_action": {
        "summary": "",
        "action": ""
    },
    "assignment_enrichment": {
        "assignment": "",
        "instructions": "",
        "guide_questions": [""]
    },
    "references": [""],
    "prepared_by": "",
    "prepared_by_title": "",
    "checked_by": "",
    "checked_by_title": "",
    "noted_by": "",
    "noted_by_title": ""
}


def build_prompt(cm_text: str, up_text: str, curriculum_topic: str, specific_lesson_focus: str, session: str, lesson_date: str, language: str, subject: str, grade_sections: str, term: str, customization: str, summative_type: str) -> str:
    return f"""
You are a Daily Learning Plan generator for teachers.

Use ONLY the supplied Curriculum Map and Unit Plan as the authoritative basis for
school-specific and lesson-specific content.

CURRICULUM MAP
==============
{cm_text}

UNIT PLAN
=========
{up_text}

REQUESTED LESSON
================
Language: {language}
Subject: {subject}
Grade and Section/s: {grade_sections}
Term: {term}
Topic from Curriculum Map: {curriculum_topic}
Specific Lesson Topic / Focus: {specific_lesson_focus}
Session: {session or "[not supplied]"}
Date: {lesson_date or "[not supplied]"}
Teacher Customization / Contextualization: {customization or "[none supplied]"}
Preferred Summative Assessment: {summative_type}

RULES
=====
1. Generate ONE realistic daily lesson plan for the Specific Lesson Topic / Focus in the selected Language (English or Filipino).
2. Use simple, natural, human teacher wording. Write student-facing content the way a teacher would actually say it in class: short, clear, warm, direct, and easy to understand on the first reading.
3. Respect Customization / Contextualization unless it conflicts with the Curriculum Map, Unit Plan, or competency.
4. Keep the COMPLETE lesson realistic for LESS THAN ONE HOUR.
5. Use ONLY Learning Competency/Competencies supported by the Curriculum Map or Unit Plan. NEVER invent one. If none matches, return an empty learning_competencies list.
6. Strictly align: Learning Competency -> Specific Lesson Focus -> Daily Objectives -> Presentation -> Activities -> Formative Assessment -> Summative Assessment -> Summary -> Action.
7. Objectives are based on the competency and MUST NOT exceed its cognitive/performance demand. If it says "discuss," do not require "demonstrate" unless supported.
8. Transfer Goal, Essential Understanding, and Essential Question MUST come from the relevant Unit Plan when available.
9. Review uses the immediately previous lesson in the source sequence and contains EXACTLY TWO simple questions, moving from basic recall/understanding to deeper thinking/application. Do not display LOTS/HOTS labels.
10. Focus is today's Specific Lesson Topic / Focus.
11. Resources use source-listed resources first; add only genuinely needed practical resources such as TV, PPT Presentation, HDMI, Textbook, Notebook, Paper, and Pen.
12. Motivation directly connects to the lesson, catches attention, is engaging, and brief.
13. Activating Prior Knowledge is a simple starter QUESTION needed for today's lesson.
14. Presentation of Concept supplies measurable daily objectives introduced in Word by "The students will be able to…". Each objective must start with the action verb only; never repeat that introductory phrase inside the objective.
15. Activities align with competency, objectives, Curriculum Map, and focus. Prefer source activities when suitable; otherwise create a simple aligned activity. Use group/individual/both only when appropriate and feasible under one hour.
16. Broadening provides one simple lesson-related question for EACH Leading, Exploring, Connecting, and Essential Question, progressively deepening thinking.
17. Ignacian Core Value is Faith, Excellence, or Service. Follow source if stated; otherwise choose the natural fit.
18. Related values: FAITH—Strong Faith in God, Prophetic Witness to Gospel Values, Nationalism, Justice, Communion. EXCELLENCE—Integrity, Competence, Resourcefulness, Discipline, Self-reliance. SERVICE—Stewardship, Humility, Charity, Courage, Preferential Love of the Poor.
19. SOCIAL ORIENTATION: connect today's lesson to family, school, community, or society. The "name" should identify the context (for example: Community, Family, School, Society). The "question_or_connection" MUST be ONE clear STUDENT-FACING QUESTION that helps learners relate or apply the lesson to that real-life social context. It must end with a question mark.
20. LESSON ACROSS DISCIPLINE: select exactly one DIFFERENT subject from: TLE, Computer, Science, Araling Panlipunan, Christian Living, Mathematics, English, Filipino. The "name" is the chosen subject. The "question_or_connection" MUST be ONE clear STUDENT-FACING QUESTION showing how today's lesson connects to that other subject. It must end with a question mark.
21. Biblical Text/Reflection genuinely aligns with the lesson. Prefer source material and NEVER fabricate verse wording.
22. Formative Assessment happens DURING the lesson/activity, checks the competency/objectives, and includes the actual ready-to-use questions/tasks/items—not only the assessment name.
23. Summative Assessment checks today's learning near the end and measures the SAME competency, objectives, and specific lesson focus.
23A. Follow the teacher's Preferred Summative Assessment when one is selected. If "Automatic - Let AI Decide" is selected, choose the most suitable assessment based on the Curriculum Map, Unit Plan, competency, topic, and available lesson time.
23B. Required item counts:
- Multiple Choice Quiz = EXACTLY 10 items. Every item must have a question, choices A-D, and the correct answer.
- Enumeration = EXACTLY 10 numbered items/prompts, each with the correct expected answer/s.
- Open-Ended / Essay = EXACTLY 3 meaningful questions, each with a concise suggested answer/key points or scoring guide.
- Long Quiz = EXACTLY 20 items. Use an appropriate aligned format or sensible mix of item types; include the correct answer for every objectively scored item and a concise key/scoring guide for open-ended items.
23C. Do not duplicate numbering inside assessment item text. The Word builder adds 1.), 2.), 3.), etc. Therefore each JSON item must begin directly with the question/task, not with a number.
23D. For multiple choice, format each JSON item as: "Question text\nA. Choice\nB. Choice\nC. Choice\nD. Choice\nAnswer: Correct choice". Keep choices plausible, clear, and aligned.
24. Summary is a QUESTION leading students to summarize/explain/apply the main learning.
25. Action is a QUESTION leading students to apply learning in a real situation.
26. Purposive Assignment/Enrichment strengthens today's learning or prepares tomorrow's lesson.
27. References: DO NOT list the Curriculum Map or Unit Plan themselves as references. Use any actual textbook/book/source named inside those documents when bibliographic details are available. Format textbook/book references in APA style as closely as the available details allow. Internet references will be added separately using Google Search grounding, so do not invent website URLs.
28. Do not invent staff names, unsupported curriculum requirements, or school-specific facts.
29. Use teacher-supplied Subject, Grade and Section/s, Term, Session, and Date where appropriate.
30. Return ONLY valid JSON.

STUDENT-FRIENDLY AND HUMANIZED WORDING
31. Write as a real classroom teacher speaking to Grade-level students, not as an academic writer, curriculum specialist, or AI.
32. Use familiar, everyday words whenever possible. Prefer short sentences and direct instructions.
33. Student-facing questions, directions, activities, motivation, review, summary, and action must be easy for students to understand on the first reading.
34. Avoid unnecessarily difficult words, abstract phrasing, jargon, long explanations, and robotic expressions.
35. Do not make simple ideas sound complicated. For example, prefer "What do you notice about the data?" over "What observations can be derived from the presented dataset?"
36. Prefer natural directions such as "Work with your group. Look at the data and answer the questions." instead of formal wording such as "Collaboratively analyze the provided data set and formulate responses."
37. Keep questions focused on ONE clear idea whenever possible. Do not combine several difficult questions into one sentence.
38. Match the wording to the students' Grade and Section/s. The thinking may be LOTS or HOTS, but the LANGUAGE of the question must still be simple and age-appropriate.
39. HOTS means deeper thinking, not harder vocabulary. Use simple words even for analysis, application, evaluation, and reflection questions.
40. For Filipino output, use natural classroom Filipino that students commonly understand. Avoid deep, old-fashioned, overly formal, or awkward literal translations. Keep official school section headings in Filipino, but make student-facing content conversational and clear.
41. For English output, use natural classroom English with common words and concise sentences.
42. Preserve the exact Learning Competency wording from the Curriculum Map/Unit Plan when it is quoted or copied. Do not simplify the official competency itself if doing so would change its meaning. Simplify the objectives, directions, questions, activities, assessments, and explanations around it.
43. Make Motivation sound inviting and interesting to students. It should feel like a short classroom hook, not a formal lesson-plan description.
44. Make Activating Prior Knowledge, Broadening questions, Summary, and Action sound like questions a teacher could naturally ask aloud in class.
45. Make assessment directions specific and simple: clearly tell students what to do, without unnecessary explanation.
46. Before returning the JSON, silently reread every student-facing sentence and simplify any wording that sounds too formal, robotic, vague, or difficult while keeping the intended learning level and curriculum alignment unchanged.
47. PRESENTATION OF CONCEPT: the Word template already prints "The students will be able to…". Therefore, each objective in presentation_of_concept MUST begin directly with the action verb (for example: "explain...", "discuss...", "identify...", "solve..."). NEVER repeat "The students will be able to" inside an objective.
48. REVIEW: identify the lesson immediately BEFORE today's lesson from the Curriculum Map/Unit Plan sequence. Write exactly TWO simple review questions about that previous lesson, progressing from recall/understanding to deeper thinking/application. Do NOT write or mention the labels LOTS or HOTS in the questions or output.
49. ASSESSMENTS MUST BE READY TO USE, not merely names such as "Oral Recitation," "Written Quiz," or "Observation." For formative assessment, provide the actual questions, tasks, prompts, computations, statements, or performance checks students will answer/do during the lesson. For summative assessment, provide the actual end-of-lesson items/tasks that directly measure today's objectives and competency. Keep them realistic for a lesson under one hour.
50. When an assessment is a quiz, include the actual quiz items. When it is a performance task, include the exact task/instructions and concise criteria needed to measure learning. Do not leave assessment content implied.
51. LESSON ACROSS DISCIPLINE: choose ONLY ONE from this allowed list: TLE, Computer, Science, Araling Panlipunan, Christian Living, Mathematics, English, Filipino. It MUST be different from the current Subject. After naming that subject, write ONE simple student-facing QUESTION that lets students integrate today's lesson with the other subject. Example format: Lesson Across Discipline: Science — "How do scientists collect and organize experimental data?" Never choose the same subject as the lesson.
51A. SOCIAL ORIENTATION: after naming the social context, write ONE simple student-facing QUESTION that asks students to connect today's lesson to family, school, community, or society. Do not write only a statement or explanation.
52. REFERENCES: prioritize references already named in the Curriculum Map or Unit Plan. Format bibliographic references in APA style as far as the available details allow. If a reliable source URL is explicitly available in the uploaded documents, include it. Never invent an author, title, year, publisher, DOI, or URL. If source details are incomplete, include only the details actually supported rather than fabricating missing information.
53. If outside information is used to enrich the lesson, identify a real, relevant, reliable source in the references. Do not output a made-up link. If no verified URL is available in the provided source context, an APA-style source without a fabricated URL is better than an invented link.
54. SIGNATURES: "prepared_by" will be replaced by the Teacher's Name entered in the generator, so do not invent a teacher name. For "checked_by", "checked_by_title", "noted_by", and "noted_by_title", copy the exact name and title found in the Curriculum Map/Unit Plan when explicitly present. Never invent or paraphrase these names/titles. If they are not found, leave those JSON fields blank so the Word template defaults can be used.
55. Keep the exact Word template structure, columns, section order, and formatting unchanged. These content rules must not alter the document layout.

RETURN EXACTLY THIS JSON SHAPE
==============================
{json.dumps(SCHEMA, indent=2)}
"""


def extract_json(text: str) -> Dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start:end+1])
        raise


def generate_plan(
    cm_text: str, up_text: str, curriculum_topic: str, specific_lesson_focus: str,
    session: str, lesson_date: str, language: str, subject: str,
    grade_sections: str, term: str, customization: str, summative_type: str,
) -> Dict[str, Any]:
    if not GEMINI_API_KEY:
        raise RuntimeError(
            "The website owner has not configured GEMINI_API_KEY in the server secrets."
        )

    client = genai.Client(api_key=GEMINI_API_KEY)
    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=build_prompt(cm_text, up_text, curriculum_topic, specific_lesson_focus, session, lesson_date, language, subject, grade_sections, term, customization, summative_type),
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            temperature=0.2,
            max_output_tokens=12000,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response. Please try again.")

    data = extract_json(response.text)
    data["references"] = clean_and_enrich_references(
        client=client,
        existing=data.get("references", []) or [],
        subject=subject,
        grade_sections=grade_sections,
        topic=curriculum_topic,
        focus=specific_lesson_focus,
        language=language,
    )
    return data


def clean_and_enrich_references(client, existing, subject, grade_sections, topic, focus, language):
    """Keep real textbook/source citations and add grounded web references when available."""
    cleaned = []
    for ref in existing:
        text = str(ref or "").strip()
        low = text.lower()
        if not text:
            continue
        if ("curriculum map" in low or "unit plan" in low) and not any(
            marker in low for marker in ["http://", "https://", "www."]
        ):
            continue
        cleaned.append(text)

    # Google Search grounding is supported by Gemini 2.5 Flash/Flash-Lite.
    # If the project's model/quota does not allow it, the lesson still generates
    # and simply keeps the verified textbook/source references from the documents.
    try:
        ref_prompt = f"""
Find 2 reliable public internet references directly related to this lesson.

Subject: {subject}
Grade/Section: {grade_sections}
Curriculum Topic: {topic}
Specific Lesson Focus: {focus}
Language of lesson: {language}

Requirements:
- Use Google Search to verify that each source is real and relevant.
- Prefer official education/government/university/professional/established educational sources.
- Do not use the school's Curriculum Map or Unit Plan as a reference.
- Return ONLY 2 lines.
- Each line must be a concise APA-style web reference followed by the full working URL.
- Never invent a URL, author, title, or date.
"""
        grounded = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=ref_prompt,
            config=types.GenerateContentConfig(
                tools=[types.Tool(google_search=types.GoogleSearch())],
                temperature=0.1,
                max_output_tokens=1200,
            ),
        )
        if grounded.text:
            for line in grounded.text.splitlines():
                line = re.sub(r"^\s*[-•*\d.)]+\s*", "", line).strip()
                if line and ("http://" in line or "https://" in line):
                    cleaned.append(line)
    except Exception:
        pass

    # Remove exact duplicates while preserving order.
    unique = []
    seen = set()
    for ref in cleaned:
        key = ref.casefold()
        if key not in seen:
            seen.add(key)
            unique.append(ref)
    return unique[:6]


# ============================================================
# WORD DOCUMENT BUILDER
# ============================================================

def set_font(run, size=10.5, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_text(p, text, *, bold=False, italic=False, size=10.5):
    r = p.add_run(text or "")
    set_font(r, size=size, bold=bold, italic=italic)
    return r


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, text, bold=True, size=11)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.18)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    add_text(p, text, bold=True)
    return p


def line(doc, label="", value="", indent=.35):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_after = Pt(2)
    if label:
        add_text(p, label, bold=True)
    if value:
        add_text(p, value)
    return p


def bullets(doc, items, indent=.48):
    for item in items or []:
        if not item:
            continue
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_after = Pt(1)
        add_text(p, item)


def numbered(doc, items, indent=.55):
    for i, item in enumerate(items or [], 1):
        if not item:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.first_line_indent = Inches(-.2)
        p.paragraph_format.space_after = Pt(1)
        add_text(p, f"{i}.) {item}")


def _school_set_run_font(run, *, bold=None, italic=None):
    """Apply the school's required Word font."""
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _school_clear_cell(cell):
    """Clear only the contents while preserving the school's table borders/size."""
    cell._tc.clear_content()


def _school_add_p(
    cell,
    text="",
    *,
    bold=False,
    italic=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left=0,
):
    p = cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if left:
        pf.left_indent = Inches(left)

    r = p.add_run(str(text or ""))
    _school_set_run_font(r, bold=bold, italic=italic)
    return p


def _school_add_label_value(cell, label, value, *, left=.35, italic_value=False):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if left:
        pf.left_indent = Inches(left)

    r = p.add_run(label)
    _school_set_run_font(r, bold=True)
    r = p.add_run(str(value or ""))
    _school_set_run_font(r, italic=italic_value)
    return p


def build_docx(d: Dict[str, Any]) -> bytes:
    """Populate the exact SMAJ sample Learning Plan format supplied by the user."""
    if not TEMPLATE_PATH.exists():
        raise RuntimeError("school_learning_plan_template.docx is missing beside app.py.")

    doc = Document(str(TEMPLATE_PATH))

    # Keep the sample's exact A4 page setup and 0.5-inch margins.
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Inches(.5)
    sec.bottom_margin = Inches(.5)
    sec.left_margin = Inches(.5)
    sec.right_margin = Inches(.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(12)

    lang = str(d.get("language", "English") or "English").strip().lower()
    fil = lang == "filipino"

    L = {
        "plan": "BANGHAY-ARALIN" if fil else "LEARNING PLAN",
        "subject": "Asignatura:" if fil else "Subject:",
        "term": "Termino:" if fil else "Term:",
        "session": "Sesyon:" if fil else "Session:",
        "topic": "Paksa:" if fil else "Topic:",
        "unit": "Yunit:" if fil else "Unit:",
        "date": "Petsa:" if fil else "Date:",
        "competencies": "Mga Kasanayang Pampagkatuto:" if fil else "Learning Competencies:",
        "transfer": "Layunin sa Paglilipat ng Pagkatuto" if fil else "Transfer Goal",
        "understanding": "Mahalagang Pag-unawa" if fil else "Essential Understanding",
        "essential_q": "Mahalagang Tanong" if fil else "Essential Question",
        "prelim": "I.  Panimulang Gawain" if fil else "I.  Preliminaries",
        "review": "A.  Balik-Aral" if fil else "A.  Review",
        "focus": "B.  Pokus" if fil else "B.  Focus",
        "resources": "C.  Mga Kagamitan" if fil else "C.  Resources",
        "motivation": "D.  Pagganyak" if fil else "D.  Motivation",
        "instruction": "Panuto: " if fil else "Instruction: ",
        "guide_q": "Gabay na Tanong:" if fil else "Guide Question:",
        "guide_qs": "Mga Gabay na Tanong:" if fil else "Guide Questions:",
        "prior": "E.  Pagpapagana ng Dating Kaalaman" if fil else "E.  Activating Prior Knowledge",
        "lesson_dev": "II.  Paglinang ng Aralin" if fil else "II.  Lesson Development",
        "presentation": "A.  Paglalahad ng Konsepto" if fil else "A.  Presentation of Concept",
        "students": "Ang mga mag-aaral ay inaasahang…" if fil else "The students will be able to…",
        "activities": "B.  Mga Gawain" if fil else "B.  Activities",
        "activity": "Gawain" if fil else "Activity",
        "individual": "Indibidwal" if fil else "Individual",
        "group": "Pangkat" if fil else "Group",
        "procedure": "Pamamaraan:" if fil else "Procedure:",
        "broadening": "C.  Pagpapalawak ng Konsepto" if fil else "C.  Broadening of Concept",
        "leading": "Panimulang Tanong" if fil else "Leading Question",
        "exploring": "Mapanuring Tanong" if fil else "Exploring Question",
        "connecting": "Tanong na Nag-uugnay" if fil else "Connecting Question",
        "essential": "Mahalagang Tanong" if fil else "Essential Question",
        "integration": "D.  Integrasyon" if fil else "D.  Integration",
        "ignacian": "Pangunahing Pagpapahalagang Ignacian" if fil else "Ignacian Core Value",
        "related": "Kaugnay na Pagpapahalaga" if fil else "Related Core Value",
        "social": "Oryentasyong Panlipunan" if fil else "Social Orientation",
        "discipline": "Ugnayan sa Ibang Disiplina" if fil else "Lesson Across Discipline",
        "biblical": "Tekstong Biblikal/Pagninilay" if fil else "Biblical Text/Reflection",
        "evaluation": "III.  Pagtataya" if fil else "III.  Evaluation/Assessment",
        "formative": "A.  Pormatibong Pagtataya" if fil else "A.  Formative Assessment",
        "summative": "B.  Sumatibong Pagtataya" if fil else "B.  Summative Assessment",
        "summary_action": "IV.  Paglalagom/Pagkilos" if fil else "IV.  Summary/Action",
        "summary": "A.  Paglalagom" if fil else "A.  Summary",
        "action": "B.  Pagkilos" if fil else "B.  Action",
        "assignment": "V.  Makabuluhang Takdang-Aralin/Pagpapayaman" if fil else "V.  Purposive Assignment/Enrichment",
        "assignment_label": "Takdang-Aralin: " if fil else "Assignment: ",
        "references": "VI.  Mga Sanggunian" if fil else "VI.  References",
        "prepared": "Inihanda ni:" if fil else "Prepared by:",
        "checked": "Sinuri ni:" if fil else "Checked by:",
        "noted": "Pinagtibay ni:" if fil else "Noted by:",
        "teacher": "Guro" if fil else "Teacher",
        "status": "Kalagayan ng Pagpapatupad" if fil else "Status of Implementation",
        "implemented": "___Naipatupad         ___ Bahagyang Naipatupad" if fil else "___Implemented         ___ Partially Implemented",
        "not_impl": "___Hindi Naipatupad" if fil else "___Not Implemented",
        "remarks": "Mga Tala:" if fil else "Remarks:",
        "observed": "Inobserbahan ni: _________________________" if fil else "Observed by: _________________________",
        "date_observed": "Petsa ng Obserbasyon: ____________________" if fil else "Date Observed: _______________________",
        "modifications": "Mga Pagbabago" if fil else "Modifications",
    }

    def clear(cell):
        _school_clear_cell(cell)

    def add(cell, text="", *, bold=False, italic=False, left=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        return _school_add_p(cell, text, bold=bold, italic=italic, left=left, align=align)

    def add_lv(cell, label, value, *, left=.35, italic_value=False):
        return _school_add_label_value(cell, label, value, left=left, italic_value=italic_value)

    # Preserve school logo and header. Translate only the plan title when Filipino is selected.
    for p in doc.paragraphs:
        if p.text.strip().upper() == "LEARNING PLAN":
            for r in p.runs:
                r.text = ""
            r = p.add_run(L["plan"])
            _school_set_run_font(r, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    # ---------------- EXACT 3 x 6 TOP INFORMATION TABLE ----------------
    meta = doc.tables[0]
    grade_section = str(d.get("grade_level", "") or "").strip()
    grade_match = re.search(r"(?:Grade\s*)?(\d+)", grade_section, flags=re.I)
    level_text = grade_match.group(1) if grade_match else ""
    subject_text = str(d.get("subject", "") or "").strip()
    subject_level = f"{subject_text} {level_text}".strip()

    row0 = [(L["subject"], subject_level), (L["term"], d.get("term", "")), (L["session"], d.get("session", ""))]
    for i, (label, value) in enumerate(row0):
        lc = meta.rows[0].cells[i*2]; vc = meta.rows[0].cells[i*2+1]
        clear(lc); clear(vc)
        add(lc, label, bold=True)
        add(vc, value)

    for i, (label, value) in enumerate([(L["topic"], d.get("topic", "")), (L["unit"], d.get("unit", ""))]):
        lc = meta.rows[1].cells[i*2]; vc = meta.rows[1].cells[i*2+1]
        clear(lc); clear(vc)
        add(lc, label, bold=True)
        add(vc, value, align=WD_ALIGN_PARAGRAPH.LEFT)

    date_cell = meta.rows[1].cells[4]
    clear(date_cell)
    add(date_cell, L["date"], bold=True)
    date_value = str(d.get("date", "") or "").strip()
    if date_value:
        for line in date_value.splitlines():
            if line.strip(): add(date_cell, f"•  {line.strip()}", left=.25, align=WD_ALIGN_PARAGRAPH.LEFT)

    comp_cell = meta.rows[2].cells[0]
    clear(comp_cell)
    add(comp_cell, L["competencies"], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, comp in enumerate(d.get("learning_competencies", []) or []):
        if str(comp).strip():
            prefix = chr(97+i) if i < 26 else str(i+1)
            add(comp_cell, f"{prefix}.  {str(comp).strip()}", left=.25, align=WD_ALIGN_PARAGRAPH.LEFT)

    # ---------------- EXACT 12-ROW MAIN TABLE ----------------
    table = doc.tables[1]
    for ri, heading, key in [(0,L["transfer"],"transfer_goal"),(1,L["understanding"],"essential_understanding"),(2,L["essential_q"],"essential_question")]:
        c=table.rows[ri].cells[0]; clear(c)
        add(c, heading, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        add(c, d.get(key,""), align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # I. Preliminaries
    prelim=d.get("preliminaries",{}) or {}
    c=table.rows[3].cells[0]; clear(c)
    add(c,L["prelim"],bold=True,left=.02,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["review"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    review=prelim.get("review",{}) or {}
    if isinstance(review,dict):
        lq=str(review.get("lots_question","") or "").strip(); hq=str(review.get("hots_question","") or "").strip()
        if lq: add(c,lq,left=.55)
        if hq: add(c,hq,left=.55)
    elif str(review).strip(): add(c,str(review).strip(),left=.55)

    add(c,L["focus"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,d.get("topic", prelim.get("focus","")),left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["resources"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    for resource in prelim.get("resources",[]) or []:
        if str(resource).strip(): add(c,f"•  {str(resource).strip()}",left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)

    add(c,L["motivation"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    mot=prelim.get("motivation",{}) or {}
    if mot.get("title"): add(c,mot.get("title",""),bold=True,italic=True,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
    if mot.get("instruction"): add_lv(c,L["instruction"],mot.get("instruction",""),left=.55)
    if mot.get("content"): add(c,mot.get("content",""),left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
    mqs=mot.get("guide_questions",[]) or []
    if mqs:
        add(c,L["guide_q"] if len(mqs)==1 else L["guide_qs"],bold=True,italic=True,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
        for i,q in enumerate(mqs,1): add(c,("" if len(mqs)==1 else f"{i}.) ")+str(q),left=.55)

    add(c,L["prior"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    prior=prelim.get("activating_prior_knowledge","")
    if isinstance(prior,list):
        for q in prior:
            if str(q).strip(): add(c,str(q).strip(),left=.55)
    elif str(prior).strip(): add(c,str(prior).strip(),left=.55)

    # II. Lesson Development
    lesson=d.get("lesson_development",{}) or {}
    c=table.rows[4].cells[0]; clear(c)
    add(c,L["lesson_dev"],bold=True,left=.02,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["presentation"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["students"],left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
    for i,obj in enumerate(lesson.get("presentation_of_concept",[]) or []):
        if str(obj).strip():
            prefix=chr(97+i) if i<26 else str(i+1)
            obj_text=str(obj).strip()
            obj_text=re.sub(r"^the\s+students?\s+will\s+be\s+able\s+to\s*[:\-–—]?\s*", "", obj_text, flags=re.I)
            obj_text=re.sub(r"^students?\s+will\s+be\s+able\s+to\s*[:\-–—]?\s*", "", obj_text, flags=re.I)
            add(c,f"{prefix}.  {obj_text}",left=.55)

    add(c,L["activities"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    for act in lesson.get("activities",[]) or []:
        typ=str(act.get("type","") or "").strip().lower()
        if "group" in typ or "pangkat" in typ: type_label=L["group"]
        elif "individual" in typ or "indibid" in typ: type_label=L["individual"]
        else: type_label=str(act.get("type","") or "").strip() or L["activity"]
        title=str(act.get("title","") or "").strip()
        add(c,f"{type_label} {L['activity']}: {title}" if title else f"{type_label} {L['activity']}",bold=True,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
        if act.get("instruction"): add_lv(c,L["instruction"],act.get("instruction",""),left=.55)
        steps=act.get("steps",[]) or []
        if steps:
            add(c,L["procedure"],bold=True,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
            for i,step in enumerate(steps,1): add(c,f"{i}.) {step}",left=.55)
        aqs=act.get("guide_questions",[]) or []
        if aqs:
            add(c,L["guide_q"] if len(aqs)==1 else L["guide_qs"],bold=True,italic=True,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
            for i,q in enumerate(aqs,1): add(c,f"{i}.) {q}",left=.55)

    add(c,L["broadening"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    broad=lesson.get("broadening_of_concept",{}) or {}
    for label,key in [(L["leading"],"leading_question"),(L["exploring"],"exploring_question"),(L["connecting"],"connecting_question"),(L["essential"],"essential_question")]:
        if broad.get(key):
            add(c,f"•  {label}",bold=True,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
            add(c,broad.get(key,""),left=.55)

    add(c,L["integration"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    integ=lesson.get("integration",{}) or {}
    for label,key,detail in [(L["ignacian"],"ignacian_core_value","connection"),(L["related"],"related_core_value","connection"),(L["social"],"social_orientation","question_or_connection"),(L["discipline"],"lesson_across_discipline","question_or_connection")]:
        item=integ.get(key,{}) or {}
        if item.get("name"):
            add_lv(c,f"•  {label}: ",item.get("name",""),left=.55)
            if item.get(detail):
                q_or_connection = str(item.get(detail,"") or "").strip()
                # Social Orientation and Lesson Across Discipline should appear as clear
                # student-facing questions on their own line.
                add(c,q_or_connection,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
    bib=integ.get("biblical_text_reflection",{}) or {}
    if bib.get("reference"): add_lv(c,f"•  {L['biblical']}: ",bib.get("reference",""),left=.55)
    if bib.get("text"): add(c,bib.get("text",""),left=.55)
    if bib.get("reflection"): add(c,bib.get("reflection",""),left=.55)

    # III. Evaluation/Assessment
    ev=d.get("evaluation_assessment",{}) or {}
    c=table.rows[5].cells[0]; clear(c)
    add(c,L["evaluation"],bold=True,left=.02,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["formative"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    for item in ev.get("formative",[]) or []:
        name=str(item.get("name","") or "").strip(); ins=str(item.get("instruction","") or "").strip()
        if name: add(c,"•  "+name,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
        if ins: add(c,ins,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
        items=item.get("items",[]) or item.get("questions",[]) or []
        for i,q in enumerate(items,1):
            qtext = re.sub(r"^\s*\d+\s*[.)]\s*", "", str(q or "").strip())
            if qtext: add(c,f"{i}.) {qtext}",left=.70,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["summative"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
    for item in ev.get("summative",[]) or []:
        name=str(item.get("name","") or "").strip(); ins=str(item.get("instruction","") or "").strip()
        if name: add(c,"•  "+name,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
        if ins: add(c,ins,left=.55,align=WD_ALIGN_PARAGRAPH.LEFT)
        items=item.get("items",[]) or item.get("questions",[]) or []
        for i,q in enumerate(items,1):
            qtext = re.sub(r"^\s*\d+\s*[.)]\s*", "", str(q or "").strip())
            if qtext: add(c,f"{i}.) {qtext}",left=.70,align=WD_ALIGN_PARAGRAPH.LEFT)

    # IV. Summary/Action
    sa=d.get("summary_action",{}) or {}
    c=table.rows[6].cells[0]; clear(c)
    add(c,L["summary_action"],bold=True,left=.02,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,L["summary"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT); add(c,sa.get("summary",""),left=.55)
    add(c,L["action"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT); add(c,sa.get("action",""),left=.55)

    # V. Assignment/Enrichment
    ass=d.get("assignment_enrichment",{}) or {}
    c=table.rows[7].cells[0]; clear(c)
    add(c,L["assignment"],bold=True,left=.02,align=WD_ALIGN_PARAGRAPH.LEFT)
    if ass.get("assignment"): add_lv(c,L["assignment_label"],ass.get("assignment",""),left=.30)
    if ass.get("instructions"): add_lv(c,L["instruction"],ass.get("instructions",""),left=.30)
    qs=ass.get("guide_questions",[]) or []
    if qs:
        add(c,L["guide_q"] if len(qs)==1 else L["guide_qs"],bold=True,left=.30,align=WD_ALIGN_PARAGRAPH.LEFT)
        for i,q in enumerate(qs,1): add(c,f"{i}.) {q}",left=.55)

    # VI. References
    c=table.rows[8].cells[0]; clear(c)
    add(c,L["references"],bold=True,left=.02,align=WD_ALIGN_PARAGRAPH.LEFT)
    for ref in d.get("references",[]) or []:
        if str(ref).strip(): add(c,str(ref).strip(),left=.30)

    # Signatures: exact sample columns.
    # Prepared by always uses the Teacher's Name entered in the generator.
    # Checked/Noted use exact source-document values when found; otherwise
    # preserve the exact names/titles from the approved school template.
    prepared_title=str(d.get("prepared_by_title","") or "").strip()
    if not prepared_title:
        prepared_title=f"{subject_text} {L['teacher']}".strip() if subject_text else L["teacher"]

    checked_name=str(d.get("checked_by","") or "").strip() or "ARIEL C. BARROS, LPT, MAEd – EM"
    checked_title=str(d.get("checked_by_title","") or "").strip() or "Academic Coordinator"
    noted_name=str(d.get("noted_by","") or "").strip() or "S. MA. LILIBETH E. MONTECLARO, RVM"
    noted_title=str(d.get("noted_by_title","") or "").strip() or "School Principal"

    for cell,label,name,title in [
        (table.rows[9].cells[0],L["prepared"],d.get("prepared_by",""),prepared_title),
        (table.rows[9].cells[2],L["checked"],checked_name,checked_title),
    ]:
        clear(cell); add(cell,label,bold=True,align=WD_ALIGN_PARAGRAPH.LEFT); add(cell,"",align=WD_ALIGN_PARAGRAPH.LEFT)
        add(cell,name,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER); add(cell,title,align=WD_ALIGN_PARAGRAPH.CENTER)

    c=table.rows[10].cells[1]; clear(c)
    add(c,L["noted"],bold=True,align=WD_ALIGN_PARAGRAPH.LEFT); add(c,"",align=WD_ALIGN_PARAGRAPH.LEFT)
    add(c,noted_name,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER); add(c,noted_title,align=WD_ALIGN_PARAGRAPH.CENTER)

    status=table.rows[11].cells[0]; mods=table.rows[11].cells[2]
    clear(status); add(status,L["status"],bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    add(status,L["implemented"],left=.05,align=WD_ALIGN_PARAGRAPH.LEFT); add(status,L["not_impl"],left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(status,L["remarks"],bold=True,left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(status,"_____________________________________",left=.05,align=WD_ALIGN_PARAGRAPH.LEFT); add(status,"_____________________________________",left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(status,"",align=WD_ALIGN_PARAGRAPH.LEFT); add(status,L["observed"],left=.05,align=WD_ALIGN_PARAGRAPH.LEFT); add(status,"",align=WD_ALIGN_PARAGRAPH.LEFT); add(status,L["date_observed"],left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)
    clear(mods); add(mods,L["modifications"],bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3): add(mods,"_____________________________________",left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)
    add(mods,"",align=WD_ALIGN_PARAGRAPH.LEFT); add(mods,L["remarks"],bold=True,left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)
    for _ in range(3): add(mods,"_____________________________________",left=.05,align=WD_ALIGN_PARAGRAPH.LEFT)

    # Ensure all generated/existing text follows the sample's Arial 12 format.
    for table_obj in doc.tables:
        for row in table_obj.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        _school_set_run_font(run, bold=run.bold, italic=run.italic)

    out=io.BytesIO(); doc.save(out); return out.getvalue()

# ============================================================
# USER INTERFACE
# ============================================================

st.subheader("1. Upload Required Documents")
c1, c2 = st.columns(2)

with c1:
    cm_file = st.file_uploader(
        " Curriculum Map (.docx)",
        type=["docx"],
        key="curriculum_map"
    )

with c2:
    up_file = st.file_uploader(
        " Unit Plan (.docx)",
        type=["docx"],
        key="unit_plan"
    )

st.subheader("2. Teacher & Lesson Details")
c_teacher, c_subject = st.columns(2)
with c_teacher:
    teacher_name = st.text_input("Teacher's Name *", placeholder="Example: Maria Santos")
with c_subject:
    subject_input = st.text_input("Subject *", placeholder="Example: Mathematics")
c_grade, c_term = st.columns(2)
with c_grade:
    grade_sections_input = st.text_input("Grade and Section/s *", placeholder="Example: Grade 8 - St. Luke, St. John")
with c_term:
    term_input = st.text_input("Term *", placeholder="Example: First Term")
language = st.selectbox("Language *", ["English", "Filipino"])
curriculum_topic = st.text_input("Topic from Curriculum Map *", placeholder="Example: Measures of Central Tendency")
specific_lesson_focus = st.text_input("Specific Lesson Topic / Focus *", placeholder="Example: Finding the Mean")
st.caption("Narrow a broad curriculum topic into the lesson for the day. The same curriculum topic may be used for several daily lesson plans.")
customization = st.text_area("Customization / Contextualization (Optional)", placeholder="Add special instructions, local examples, preferred activities, student context, or anything you want adjusted.", height=110)

summative_type = st.selectbox(
    "Preferred Summative Assessment (Optional)",
    [
        "Automatic - Let AI Decide",
        "Multiple Choice Quiz - 10 items",
        "Enumeration - 10 items",
        "Open-Ended / Essay - 3 items",
        "Long Quiz - 20 items",
    ],
    help="Optional. If you leave this on Automatic, the AI will choose the most suitable assessment based on the Curriculum Map, Unit Plan, competency, and lesson topic.",
)
st.caption("Assessment items will be aligned with the lesson. Multiple-choice items include choices and the correct answer.")

c3, c4 = st.columns(2)
with c3: session = st.text_input("Session (Optional)", placeholder="Example: Session 1")
with c4: lesson_date = st.text_input("Date (Optional)", placeholder="Example: September 15, 2026")

st.subheader("3. Requirements Check")

ready = True
for uploaded, ok_text, no_text in [(cm_file,"Curriculum Map uploaded","Curriculum Map required"),(up_file,"Unit Plan uploaded","Unit Plan required")]:
    if uploaded: st.markdown(f'<div class="status-ok">{ok_text}</div>', unsafe_allow_html=True)
    else:
        st.markdown(f'<div class="status-no">{no_text}</div>', unsafe_allow_html=True); ready=False
for label,value in [("Teacher",teacher_name),("Subject",subject_input),("Grade and Section/s",grade_sections_input),("Term",term_input),("Curriculum Topic",curriculum_topic),("Lesson Focus",specific_lesson_focus)]:
    if value.strip(): st.markdown(f'<div class="status-ok">{label}: {value.strip()}</div>', unsafe_allow_html=True)
    else:
        st.markdown(f'<div class="status-no">{label} required</div>', unsafe_allow_html=True); ready=False
st.markdown(f'<div class="status-ok">Language: {language}</div>', unsafe_allow_html=True)

if not GEMINI_API_KEY:
    st.warning(
        "Website setup is incomplete: the owner still needs to add the server API key."
    )

st.divider()

if st.button(
    " GENERATE LESSON PLAN",
    type="primary",
    disabled=(not ready or not GEMINI_API_KEY),
):
    try:
        with st.spinner("Reading and aligning your curriculum documents..."):
            cm_text = read_docx(cm_file)
            up_text = read_docx(up_file)

        with st.spinner("Creating your Daily Learning Plan..."):
            data = generate_plan(
                cm_text=cm_text, up_text=up_text, curriculum_topic=curriculum_topic.strip(),
                specific_lesson_focus=specific_lesson_focus.strip(), session=session.strip(),
                lesson_date=lesson_date.strip(), language=language, subject=subject_input.strip(),
                grade_sections=grade_sections_input.strip(), term=term_input.strip(), customization=customization.strip(),
                summative_type=summative_type,
            )
            data["language"]=language
            data["subject"]=subject_input.strip()
            data["grade_level"]=grade_sections_input.strip()
            data["term"]=term_input.strip()
            data["prepared_by"]=teacher_name.strip()
            data["topic"]=specific_lesson_focus.strip()

            competencies = data.get("learning_competencies", []) or []
            competencies = [str(c).strip() for c in competencies if str(c).strip()]
            data["learning_competencies"] = competencies

            if not competencies:
                st.warning(
                    "No matching Learning Competency was found in the uploaded Curriculum Map or Unit Plan "
                    "for this topic and lesson focus. Please check the topic/focus and try again."
                )
                st.stop()

        with st.spinner("Preparing the Word document..."):
            docx_bytes = build_docx(data)

        st.session_state["generated_docx"] = docx_bytes
        st.session_state["generated_topic"] = specific_lesson_focus.strip()
        st.success(" Lesson plan ready!")

    except Exception as exc:
        raw_error = str(exc)
        error_text = raw_error.lower()

        if "429" in error_text or "quota" in error_text or "resource_exhausted" in error_text or "rate" in error_text:
            st.error(
                "The Gemini free-tier limit has been reached for now. Please wait and try again later."
            )
        elif (
            "api key" in error_text
            or "api_key" in error_text
            or "permission" in error_text
            or "unauthenticated" in error_text
            or "401" in error_text
            or "403" in error_text
        ):
            st.error(
                "Gemini could not authenticate this app. Please check the GEMINI_API_KEY saved in Streamlit Secrets."
            )
        elif "404" in error_text or "not found" in error_text or "model" in error_text and "not" in error_text:
            st.error(
                "The configured Gemini model is unavailable. Check GEMINI_MODEL in Streamlit Secrets."
            )
        elif "json" in error_text or "decode" in error_text:
            st.error(
                "Gemini returned an incomplete lesson plan response. Please click Generate Lesson Plan again."
            )
        else:
            st.error(
                "The lesson plan could not be generated. The app owner can use the diagnostic message below."
            )
            # Show only a short diagnostic message. API keys are redacted if they ever appear.
            safe_error = re.sub(r'AIza[0-9A-Za-z_-]{20,}', '[REDACTED_API_KEY]', raw_error)
            safe_error = safe_error[:800]
            st.code(f"{type(exc).__name__}: {safe_error}")

        print(f"Lesson plan generation error ({type(exc).__name__}): {raw_error}")

if "generated_docx" in st.session_state:
    safe_topic = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        st.session_state.get("generated_topic", "Lesson")
    ).strip("_") or "Lesson"

    st.download_button(
        " DOWNLOAD WORD FILE (.DOCX)",
        data=st.session_state["generated_docx"],
        file_name=f"Daily_Learning_Plan_{safe_topic}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
    )

st.markdown(
    ''
    '<p class="small" style="text-align:center;margin-top:.3rem;">'
    'Your uploaded files are used to generate the requested lesson plan during this session.'
    '</p>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="developer">Developer: Aurie Joy D. Ellevera</div>',
    unsafe_allow_html=True
)
